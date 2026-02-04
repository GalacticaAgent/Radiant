import fitz
from docx import Document
import json
import os

from utils.llm_client import LLMClient
from utils.prompt_builder import build_persona_prompt, build_focus_prompt


class AreaChair:

    def __init__(self, chair_id: str, chair_profile: dict):

        self.chair_id = chair_id
        self.profile = chair_profile
        self.paper_content = None

        self.llm = LLMClient()

        persona_prompt = build_persona_prompt(
            self.profile.get("persona", {})
        )
        focus_prompt = build_focus_prompt(
            self.profile.get("focus", [])
        )

        self.system_prompt = f"""
            You are a senior academic mentor.
    
            {persona_prompt}
            {focus_prompt}
    
            CORE RULES:
            - Prioritize intellectual clarity over tone
            - Be precise, explicit, and demanding
            - Do NOT consider author emotions at this stage
            - Do NOT soften logical standards
            - Do NOT issue accept/reject judgments
            """.strip()

    # ============================
    # Paper loading
    # ============================

    def load_paper(self, file_path: str):

        ext = os.path.splitext(file_path)[-1].lower()

        if ext == ".pdf":
            self.paper_content = self._read_pdf(file_path)
        elif ext == ".docx":
            self.paper_content = self._read_docx(file_path)
        elif ext == ".txt":
            with open(file_path, "r", encoding="utf-8") as f:
                self.paper_content = f.read()
        else:
            raise ValueError("Unsupported file type")

        return self.paper_content

    def _read_pdf(self, path):

        doc = fitz.open(path)
        pages = []
        for page in doc:
            blocks = page.get_text("blocks")
            blocks.sort(key=lambda b: (b[1], b[0]))
            text = "\n".join(b[4] for b in blocks)
            pages.append(f"--- Page {page.number + 1} ---\n{text}")
        return "\n".join(pages)

    def _read_docx(self, path):

        doc = Document(path)
        parts = []

        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text)

        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))

        return "\n".join(parts)

    # =====================================================
    # Output
    # =====================================================
    def synthesize_feedback(self, reviewer_outputs: dict) -> dict:

        core_tensions = self._analyze_reviewer_concerns(reviewer_outputs)
        revision_guidance = self._build_revision_guidance(core_tensions)
        formatted_tensions = self._format_core_tensions(core_tensions)
        formatted_revision_map = self._format_revision_map(revision_guidance)
        final_review = self._final_review(formatted_tensions, formatted_revision_map)

        return {
            "core_tensions_raw": core_tensions,
            "revision_guidance_raw": revision_guidance,
            "core_tensions": formatted_tensions,
            "revision_map": formatted_revision_map,
            "final_review": final_review
        }

    # =====================================================
    # Stage 1: Generate core tensions
    # =====================================================

    def _analyze_reviewer_concerns(self, reviewer_outputs: dict) -> dict:

        user_prompt = f"""
            Reviewer Reports:
            {json.dumps(reviewer_outputs, indent=2)}
    
            TASK:
            Identify 2–4 CORE INTELLECTUAL TENSIONS.
    
            A core tension is:
            - A point where the paper’s central ambition meets an unmet academic standard
            - A reason reviewers continue to press despite revisions
    
            OUTPUT (STRICT JSON):
            {{
              "core_tensions": [
                {{
                  "theme": "Concise analytical label",
                  "implicit_standard": "What academic standard reviewers are enforcing",
                  "gap_description": "What is currently missing or insufficient",
                  "intellectual_consequence": "What fails if this gap remains",
                  "reviewer_sources": ["reviewer_id"]
                }}
              ]
            }}
    
            RULES:
            - Be analytically explicit
            - Do NOT moderate tone
            - Do NOT consider author feelings
            - No markdown, no extra text
            
            STRICTLY PROHIBITED: Do not include any markdown formatting like ```json or any introductory text.
            """

        response = self.llm.ask(
            system_prompt=self.system_prompt,
            user_prompt=user_prompt
        )
        return json.loads(response)

    # =====================================================
    # Stage 2: Generate revision reasoning
    # =====================================================

    def _build_revision_guidance(self, core_tensions: dict) -> dict:

        user_prompt = f"""
            Core Tensions:
            {json.dumps(core_tensions, indent=2)}
    
            TASK:
            For each tension, describe how a competent scholar
            would reason their way toward resolving it.
    
            OUTPUT (STRICT JSON):
            {{
              "revision_guidance": [
                {{
                  "tension": "Theme",
                  "conceptual_work_needed": "What needs to be reconsidered at the level of ideas",
                  "evidence_logic": "What kind of reasoning or evidence would meet the standard",
                  "insufficient_fix": "A common but inadequate response",
                  "expected_gain": "What becomes more robust if resolved"
                }}
              ]
            }}
    
            RULES:
            - Focus on reasoning, not writing
            - No pedagogical tone
            - No reassurance language
            - No markdown
            
            STRICTLY PROHIBITED: Do not include any markdown formatting like ```json or any introductory text.
            """

        response = self.llm.ask(
            system_prompt=self.system_prompt,
            user_prompt=user_prompt
        )
        return json.loads(response)

    # =====================================================
    # Stage 3: Format tensions for human reading
    # =====================================================

    def _format_core_tensions(self, raw_tensions: dict) -> dict:

        user_prompt = f"""
            Raw Core Tensions:
            {json.dumps(raw_tensions, indent=2)}
    
            TASK:
            Rewrite the core tensions for the authors.
    
            STYLE REQUIREMENTS:
            - Preserve full logical force
            - Avoid accusatory or dismissive wording
            - Frame issues as intellectual tensions, not failures
            - Use calm, professional academic language
    
            OUTPUT (STRICT JSON):
            {{
              "core_tensions": [
                {{
                  "theme": "Theme",
                  "why_this_matters": "Why reviewers care about this point",
                  "where_the_paper_struggles": "What is currently unclear or unsupported",
                  "why_it_deserves_attention": "Why resolving this would strengthen the work"
                }}
              ]
            }}
    
            No markdown. No extra text.
            STRICTLY PROHIBITED: Do not include any markdown formatting like ```json or any introductory text.
            """

        response = self.llm.ask(
            system_prompt=self.system_prompt,
            user_prompt=user_prompt
        )
        return json.loads(response)

    # =====================================================
    # Stage 4: Format revision guidance for authors
    # =====================================================

    def _format_revision_map(self, raw_guidance: dict) -> dict:

        user_prompt = f"""
            Raw Revision Guidance:
            {json.dumps(raw_guidance, indent=2)}
    
            TASK:
            Rewrite this guidance so it supports thoughtful revision
            without sounding prescriptive or corrective.
    
            STYLE:
            - Respectful and encouraging
            - Maintain intellectual rigor
            - Suggest directions of thought, not instructions
    
            OUTPUT (STRICT JSON):
            {{
              "revision_map": [
                {{
                  "tension": "Theme",
                  "reflection_focus": "What the authors might reflect on",
                  "what_reviewers_are_looking_for": "What would help readers be convinced",
                  "pitfall_to_avoid": "What may look sufficient but is not",
                  "how_the_paper_improves": "What becomes clearer or stronger"
                }}
              ]
            }}
    
            No markdown. No extra text.
            STRICTLY PROHIBITED: Do not include any markdown formatting like ```json or any introductory text.
            """

        response = self.llm.ask(
            system_prompt=self.system_prompt,
            user_prompt=user_prompt
        )
        return json.loads(response)

    # =====================================================
    # Stage 5: Write Final Review
    # =====================================================

    def _final_review(self, formatted_tensions: dict, formatted_revision_map: dict) -> str:

        user_prompt = f"""
            DATA INPUTS:
            - Core Conflicts/Tensions: {json.dumps(formatted_tensions, indent=2)}
            - Proposed Revisions: {json.dumps(formatted_revision_map, indent=2)}

            TASK:
            1. Provide a concise overall evaluation of the paper's potential and its current status.
            2. Categorize all required modifications into three levels of severity:
                - CRITICAL (Fatal Flaws): Issues that will lead to a definitive Reject if not fixed.
                - MAJOR (Required Changes): Significant issues affecting the paper's impact or validity.
                - MINOR (Polishing): Clarity, formatting, and secondary suggestions.
            3. For each point, state the ISSUE and the REASON (why this matters for the paper's logic).
            4. Rank the points by urgency within each category.

            OUTPUT FORMAT:
            ## Overall Evaluation
            [Your executive summary]

            ---

            ## Prioritized Revision Roadmap

            ### Critical - Must Address for Re-evaluation
            1. **[Short Title of Issue]**
                - **Reason**: [Why it's a fatal flaw]
                - **Action Required**: [Briefly what needs to be changed]

            ### Major - Significant Improvements Required
            1. **[Short Title of Issue]**
                - **Reason**: [The logical or experimental gap]
                - **Action Required**: [Specific steps]

            ### Minor - Presentation & Clarity
                - [List points]

            STYLE:
                - Direct, authoritative, and intellectually demanding.
                - Focus on causality and evidence.
                - Do not be vague (e.g., avoid "improve clarity", say "redefine the objective function in Eq 3").
            """

        return self.llm.ask(
            system_prompt=self.system_prompt,
            user_prompt=user_prompt
        )
